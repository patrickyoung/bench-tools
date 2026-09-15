#!/usr/bin/env python3
"""Word composition from an executive writer's structured prose and selected figures."""
import json,sys,subprocess,os,shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_BREAK
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(sys.argv[1]).resolve();s=json.loads((root/'output/spec.json').read_text())['content'];story=json.loads((root/'inputs/story.json').read_text())['content'];source=json.loads((root/'inputs/source.json').read_text());design=json.loads((root/'inputs/design.json').read_text())['content'];theme=story['design'];out=root/'output';out.mkdir(exist_ok=True);doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.8)
styles=doc.styles;normal=styles['Normal'];normal.font.name=theme['font'];normal.font.size=Pt(11);normal.font.color.rgb=RGBColor.from_string(theme['ink'][1:]);normal.paragraph_format.space_after=Pt(8);normal.paragraph_format.line_spacing=1.12
styles['Caption'].font.name=theme['font'];styles['Caption'].font.size=Pt(9);styles['Caption'].font.bold=False;styles['Caption'].font.color.rgb=RGBColor.from_string('526476')
for style in styles:
 for borders in style.element.findall('.//'+qn('w:pBdr')):borders.getparent().remove(borders)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.add_run('Page ').font.size=Pt(9);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
subprocess.run([os.environ['PUBLICATION_PLOT_PYTHON'],str(Path(__file__).with_name('render_graphics.py')),str(root),'--document'],check=True,timeout=150)
figure_notes=json.loads((root/'build/document-graphics/notes.json').read_text())
for name,size in [('Title',30),('Subtitle',15),('Heading 1',20),('Heading 2',14)]:
 styles[name].font.name=theme['font'];styles[name].font.size=Pt(size);styles[name].font.color.rgb=RGBColor.from_string(theme['ink'][1:]);styles[name].paragraph_format.space_before=Pt(14);styles[name].paragraph_format.space_after=Pt(8)
doc.add_paragraph(s['title'],'Title');doc.add_paragraph(s['subtitle'],'Subtitle');doc.add_heading('Executive summary',1)
for text in s['executive_summary'].split('\n\n'):
 p=doc.add_paragraph();label,sep,body=text.partition(': ')
 if sep and len(label)<65:p.add_run(label+': ').bold=True;p.add_run(body)
 else:p.add_run(text)
d=source['decision'];p=doc.add_paragraph();p.add_run('Decision status: ').bold=True;p.add_run(d['status'].capitalize()+(' · '+source['candidate_names'].get(d['candidate_id'],'') if d['candidate_id'] else ''))
matrix=source['matrix'];table=doc.add_table(rows=1, cols=4);table.style='Normal Table'
for c,text in zip(table.rows[0].cells,['Option','Score bounds /100','Evidence coverage','Gate eligibility']):c.text=text
header=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header)
for r in matrix['totals']:
 for c,text in zip(table.add_row().cells,[r['name'],f"{r['lower_bound']:g}–{r['upper_bound']:g}",f"{r['coverage_percent']:g}%",r['eligibility']]):c.text=text
doc.add_paragraph('Score bounds describe missing evidence. Coverage does not establish source truth or statistical certainty.','Caption')
visuals={v['id']:v for v in design['visuals']};facts={f['id']:f for f in source['facts']};markdown=['# '+s['title'],s['subtitle'],'## Executive summary',s['executive_summary']]
used=[];card_elements=set()
for part in s['sections']:
 doc.add_heading(part['heading'],1);markdown+=['## '+part['heading']]
 for para in part['paragraphs']:
  p=doc.add_paragraph();label,sep,body=para.partition(': ')
  if sep and len(label)<65:p.add_run(label+': ').bold=True;p.add_run(body)
  else:p.add_run(para)
  markdown.append(para)
 if part['visual_id']:
  v=visuals[part['visual_id']];img=root/'build/document-graphics'/(v['id']+'.png');doc.add_heading(v['title'],2)
  if v['kind']=='decision_path':
   cards=doc.add_table(rows=(len(v['steps'])+1)//2,cols=2);cards.style='Normal Table';card_elements.add(cards._tbl)
   for i,step in enumerate(v['steps']):
    cell=cards.rows[i//2].cells[i%2];cell.paragraphs[0].text=f"{i+1:02d}  {step['heading']}";cell.paragraphs[0].paragraph_format.keep_with_next=True;cell.add_paragraph(step['detail'])
   borders=OxmlElement('w:tblBorders')
   for side in ['insideH','insideV']:
    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'32');edge.set(qn('w:color'),'FFFFFF');borders.append(edge)
   cards._tbl.tblPr.append(borders)
  else:
   p=doc.add_paragraph();run=p.add_run();run.add_picture(str(img),width=Inches(6.8));desc=run._r.xpath('.//wp:docPr')
   if desc:desc[0].set('descr',v['alt'])
   p.paragraph_format.keep_with_next=True
  caption=v['subtitle']+' '+v['caption']
  doc.add_paragraph(caption,'Caption');markdown+=['![ '+v['alt']+' ](graphics/'+v['id']+'.png)',caption]
  (out/'graphics').mkdir(exist_ok=True);shutil.copyfile(root/'inputs/graphics'/img.name,out/'graphics'/img.name)
 used+=part['fact_ids']
heading=doc.add_heading('Evidence and method notes',1);heading.paragraph_format.page_break_before=True
references=doc.add_table(rows=1,cols=2);references.style='Normal Table';references.autofit=False;references.columns[0].width=Inches(2.3);references.columns[1].width=Inches(4.5)
references.rows[0].cells[0].text='Evidence or method';references.rows[0].cells[1].text='Source / location';repeat=OxmlElement('w:tblHeader');references.rows[0]._tr.get_or_add_trPr().append(repeat)
for fid in dict.fromkeys(used):
 f=facts[fid];cells=references.add_row().cells;cells[0].text=f['label'];cells[1].text=f['source'];markdown+=['- '+f['label']+': '+f['source']]
 for cell in cells:
  for p in cell.paragraphs:p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.0
for t in doc.tables:
 is_card=t._tbl in card_elements
 for index,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
  for cell in row.cells:
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDF1F4' if is_card else theme['ink'][1:] if index==0 else 'EDF1F4' if index%2 else 'FFFFFF');cell._tc.get_or_add_tcPr().append(shade)
   for pi,para in enumerate(cell.paragraphs):
    if is_card:para.paragraph_format.space_before=Pt(6);para.paragraph_format.space_after=Pt(6)
    for run in para.runs:run.font.size=Pt(12 if is_card and pi==0 else 11 if is_card else 10);run.font.color.rgb=RGBColor.from_string(theme['accent'][1:] if is_card and pi==0 else theme['ink'][1:] if is_card else 'FFFFFF' if index==0 else theme['ink'][1:]);run.bold=(pi==0) if is_card else index==0
doc.core_properties.title=s['title'][:255];doc.core_properties.author='';doc.core_properties.subject=s['subtitle'][:255];doc.save(out/'report.docx');(out/'report.md').write_text('\n\n'.join(markdown)+'\n')
preview=root/'previews/document';preview.mkdir(parents=True,exist_ok=True)
subprocess.run([sys.executable,os.environ['PUBLICATION_DOCX_RENDERER'],str(out/'report.docx'),'--output_dir',str(preview),'--emit_pdf'],check=True,timeout=180)
pdf=preview/'report.pdf'
if not pdf.is_file():raise RuntimeError('Missing rendered document PDF')
(out/'report.pdf').write_bytes(pdf.read_bytes())
